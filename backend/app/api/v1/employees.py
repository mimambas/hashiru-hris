"""Employee CRUD, import, history."""
import csv
import io
import re
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from sqlalchemy.orm import Session

from app.core.deps import (
    get_current_user,
    get_db,
    log_audit,
    paginate,
    pagination_params,
    require_role,
)
from app.models import AuditLog, Employee, User
from app.schemas import (
    SELF_EDITABLE_FIELDS,
    EmployeeCreate,
    EmployeeDetail,
    EmployeeListItem,
    EmployeeUpdate,
)

router = APIRouter()

HR_ROLES = ("super_admin", "hr_director", "hr_manager", "hr_officer")
MANAGER_ROLES = ("super_admin", "hr_director", "hr_manager", "hr_officer", "dept_manager", "team_leader")


def _nested(emp: Employee, db: Session) -> dict:
    from app.models import Department, Position

    dept = db.query(Department).filter(Department.id == emp.department_id).first() if emp.department_id else None
    pos = db.query(Position).filter(Position.id == emp.position_id).first() if emp.position_id else None
    return {
        **{c.name: getattr(emp, c.name) for c in emp.__table__.columns},
        "department": (
            {"id": dept.id, "name": dept.name, "code": dept.code, "parent_id": dept.parent_id,
             "head_id": dept.head_id, "cost_center": dept.cost_center} if dept else None
        ),
        "position": (
            {"id": pos.id, "title": pos.title, "code": pos.code, "level": pos.level,
             "grade": pos.grade, "department_id": pos.department_id,
             "min_salary": pos.min_salary, "max_salary": pos.max_salary} if pos else None
        ),
    }


def _validate_nik(nik: str | None) -> None:
    if nik and (not re.fullmatch(r"\d{16}", nik)):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "NIK harus 16 digit angka")


def _next_employee_id(db: Session) -> str:
    today = date.today().strftime("%Y%m%d")
    count = db.query(Employee).filter(Employee.employee_id.like(f"EMP-{today}-%")).count()
    return f"EMP-{today}-{count + 1:03d}"


@router.get("", response_model=dict)
def list_employees(
    q: str | None = None,
    department_id: str | None = None,
    status_: str | None = Query(None, alias="status"),
    paging: dict = Depends(pagination_params),
    db: Session = Depends(get_db),
    user: User = Depends(require_role("super_admin", "hr_director", "hr_manager", "hr_officer",
                                      "recruiter", "finance_officer", "dept_manager", "team_leader", "employee")),
):
    query = db.query(Employee).order_by(Employee.full_name)
    if q:
        like = f"%{q}%"
        query = query.filter((Employee.full_name.ilike(like)) | (Employee.employee_id.ilike(like)) | (Employee.email.ilike(like)))
    if department_id:
        query = query.filter(Employee.department_id == department_id)
    if status_:
        query = query.filter(Employee.status == status_)
    return paginate(query, paging["page"], paging["per_page"], lambda e: EmployeeListItem(**_nested(e, db)).model_dump())


@router.post("", response_model=EmployeeDetail, status_code=status.HTTP_201_CREATED)
def create_employee(
    payload: EmployeeCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_role(*HR_ROLES)),
):
    _validate_nik(payload.nik)
    if payload.nik and db.query(Employee).filter(Employee.nik == payload.nik).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "NIK sudah terdaftar")
    if payload.npwp and db.query(Employee).filter(Employee.npwp == payload.npwp).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "NPWP sudah terdaftar")
    if db.query(Employee).filter(Employee.email == payload.email).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "Email sudah terdaftar")
    emp = Employee(**payload.model_dump(), employee_id=_next_employee_id(db))
    db.add(emp)
    db.flush()
    log_audit(db, "employee", emp.id, "create", user.id, {"employee_id": emp.employee_id})
    db.commit()
    db.refresh(emp)
    return EmployeeDetail(**_nested(emp, db))


@router.get("/{emp_id}", response_model=EmployeeDetail)
def get_employee(
    emp_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Karyawan tidak ditemukan")
    if user.role == "employee" and user.employee_id != emp.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki akses ke data ini")
    return EmployeeDetail(**_nested(emp, db))


@router.put("/{emp_id}", response_model=EmployeeDetail)
def update_employee(
    emp_id: str,
    payload: EmployeeUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Karyawan tidak ditemukan")
    data = {k: v for k, v in payload.model_dump().items() if v is not None}
    if user.role == "employee":
        if user.employee_id != emp.id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki akses ke data ini")
        forbidden = set(data) - SELF_EDITABLE_FIELDS
        if forbidden:
            raise HTTPException(
                status.HTTP_403_FORBIDDEN,
                f"Field {', '.join(sorted(forbidden))} hanya dapat diubah oleh HR",
            )
    elif user.role not in ("super_admin", "hr_director", "hr_manager", "hr_officer", "recruiter"):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki akses ke fitur ini")
    if "nik" in data:
        _validate_nik(data["nik"])
        if data["nik"] and db.query(Employee).filter(Employee.nik == data["nik"], Employee.id != emp.id).first():
            raise HTTPException(status.HTTP_409_CONFLICT, "NIK sudah terdaftar")
    if "npwp" in data and data["npwp"]:
        if db.query(Employee).filter(Employee.npwp == data["npwp"], Employee.id != emp.id).first():
            raise HTTPException(status.HTTP_409_CONFLICT, "NPWP sudah terdaftar")
    changes = {k: {"old": getattr(emp, k), "new": v} for k, v in data.items() if getattr(emp, k) != v}
    for k, v in data.items():
        setattr(emp, k, v)
    log_audit(db, "employee", emp.id, "update", user.id, {k: str(v["new"]) for k, v in changes.items()})
    db.commit()
    db.refresh(emp)
    return EmployeeDetail(**_nested(emp, db))


@router.delete("/{emp_id}")
def delete_employee(
    emp_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(require_role("super_admin", "hr_manager")),
):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Karyawan tidak ditemukan")
    emp.status = "inactive"
    log_audit(db, "employee", emp.id, "soft_delete", user.id, {})
    db.commit()
    return {"ok": True, "message": "Karyawan dinonaktifkan"}


@router.post("/import")
def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_role(*HR_ROLES)),
):
    name = (file.filename or "").lower()
    try:
        raw = file.file.read()
        rows: list[dict] = []
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), read_only=True)
            ws = wb.active
            headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for r in ws.iter_rows(min_row=2, values_only=True):
                rows.append({h: ("" if v is None else str(v).strip()) for h, v in zip(headers, r)})
        else:  # CSV
            text = raw.decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text)))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"File tidak dapat dibaca: {exc}") from exc

    imported = 0
    errors: list[dict] = []
    for idx, row in enumerate(rows, start=2):
        try:
            full_name = (row.get("full_name") or "").strip()
            email = (row.get("email") or "").strip().lower()
            join_raw = (row.get("join_date") or "").strip()
            if not full_name or not email or not join_raw:
                raise ValueError("full_name, email, dan join_date wajib diisi")
            join_date = date.fromisoformat(join_raw)
            nik = (row.get("nik") or "").strip() or None
            _validate_nik(nik)
            if nik and db.query(Employee).filter(Employee.nik == nik).first():
                raise ValueError("NIK sudah terdaftar")
            if db.query(Employee).filter(Employee.email == email).first():
                raise ValueError("Email sudah terdaftar")
            base = row.get("base_salary") or "0"
            emp = Employee(
                employee_id=_next_employee_id(db),
                full_name=full_name,
                email=email,
                join_date=join_date,
                nik=nik,
                phone=(row.get("phone") or "").strip() or None,
                department_id=(row.get("department_id") or "").strip() or None,
                position_id=(row.get("position_id") or "").strip() or None,
                base_salary=base,
                ptkp_status=(row.get("ptkp_status") or "TK/0").strip(),
                employment_status=(row.get("employment_status") or "contract").strip(),
                status="active",
            )
            db.add(emp)
            imported += 1
        except (ValueError, HTTPException) as exc:
            errors.append({"row": idx, "reason": str(exc.detail if isinstance(exc, HTTPException) else exc)})
    db.commit()
    log_audit(db, "employee", "-", "bulk_import", user.id, {"imported": imported, "errors": len(errors)})
    return {"imported": imported, "errors": errors}


@router.get("/{emp_id}/history")
def employee_history(
    emp_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(require_role(*HR_ROLES, "dept_manager", "team_leader")),
):
    logs = (
        db.query(AuditLog)
        .filter(AuditLog.entity_type == "employee", AuditLog.entity_id == emp_id)
        .order_by(AuditLog.created_at.desc())
        .all()
    )
    return [
        {"id": l.id, "action": l.action, "actor_id": l.actor_id, "changes": l.changes, "created_at": l.created_at}
        for l in logs
    ]


@router.get("/{emp_id}/export")
def employee_export(
    emp_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(require_role(*HR_ROLES)),
):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Karyawan tidak ditemukan")
    return _nested(emp, db)
